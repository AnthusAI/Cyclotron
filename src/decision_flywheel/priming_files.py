"""Bounded structured-file previews and classifier-derived empty templates."""
import base64
import csv
from copy import copy
from datetime import date, datetime
import io
import json
import zipfile
from jsonschema import Draft202012Validator
from .models import DecisionTask, Item
from .workspace_catalog import encoded, timestamp

MAX_BYTES = 4_000_000
MAX_ROWS = 200


def checked_schema(schema):
    if not isinstance(schema,dict) or schema.get('type') != 'object':
        raise ValueError('item values schema must describe an object')
    def check(value):
        if isinstance(value,dict):
            if '$ref' in value and not value['$ref'].startswith('#'):
                raise ValueError('remote schema references are not permitted')
            for child in value.values():check(child)
        elif isinstance(value,list):
            for child in value:check(child)
    check(schema)
    Draft202012Validator.check_schema(schema)
    return schema


def columns(schema,classifiers):
    fields=['values.'+key for key in checked_schema(schema).get('properties',{})]
    if not fields:fields=['values']
    return ['id','occurred_at',*fields,*[f'{row["id"]}.{field}' for row in classifiers for field in ('label','rationale')]]


def template(kind,schema,classifiers):
    names=columns(schema,classifiers)
    result={'format':kind,'classifiers':[{'id':row['id'],'revision':row['revision'],'labels':[c['label'] for c in row['config']['classes']]} for row in classifiers],'schema':schema}
    if kind=='csv':
        stream=io.StringIO();csv.writer(stream).writerow(names)
        return {**result,'content':stream.getvalue(),'encoding':'utf-8'}
    if kind!='xlsx':raise ValueError('template format must be csv or xlsx')
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    wb=Workbook();ws=wb.active;ws.title='Examples';ws.append(names)
    ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:{get_column_letter(len(names))}1'
    choices=wb.create_sheet('Labels');instructions=wb.create_sheet('Instructions')
    instructions.append(['Fill Examples with schema-valid item fields. Labels are required for every selected classifier; explanations are optional.'])
    instructions.append(['Use ISO 8601 dates. Objects and arrays use JSON. Formulas are rejected. Preview all rows before acceptance.'])
    for i,row in enumerate(classifiers,1):
        labels=[c['label'] for c in row['config']['classes']]
        choices.cell(1,i,row['name']+f' · revision {row["revision"]}')
        for j,label in enumerate(labels,2):choices.cell(j,i,label).data_type='s'
        letter=get_column_letter(i)
        rule=DataValidation(type='list',formula1=f"'Labels'!${letter}$2:${letter}${len(labels)+1}",allow_blank=False)
        rule.errorTitle='Choose a configured label';rule.error='Select a label from the Labels sheet.';rule.showErrorMessage=True
        ws.add_data_validation(rule);column=get_column_letter(names.index(row['id']+'.label')+1);rule.add(f'{column}2:{column}{MAX_ROWS+1}')
    for cell in ws[1]:
        font=copy(cell.font);font.bold=True;cell.font=font
    for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=26
    stream=io.BytesIO();wb.save(stream)
    return {**result,'content':base64.b64encode(stream.getvalue()).decode(),'encoding':'base64'}


def error(row,field,message):return {'row':row,'field':field,'message':message}


def read_rows(kind,content,schema,classifiers):
    try:raw=base64.b64decode(content,validate=True)
    except Exception:raise ValueError('file must be base64 encoded') from None
    if len(raw)>MAX_BYTES:raise ValueError('file exceeds size limit')
    errors=[]
    if kind in ('json','jsonl'):
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:raise ValueError('file must be UTF-8') from None
        def parse(value):return json.loads(value,parse_constant=lambda _:(_ for _ in ()).throw(ValueError('nonfinite number')))
        if kind=='json':
            try:rows=parse(text)
            except (ValueError,TypeError):return [],[error(1,'file','Invalid JSON')]
            if not isinstance(rows,list):return [],[error(1,'file','JSON must be an array of rows')]
        else:
            rows=[]
            for n,line in enumerate(text.splitlines(),1):
                if not line.strip():continue
                try:rows.append(parse(line))
                except ValueError:rows.append(None);errors.append(error(n,'file','Invalid JSONL row'))
    elif kind in ('csv','xlsx'):
        if kind=='csv':
            try:table=list(csv.reader(io.StringIO(raw.decode('utf-8-sig'))))
            except (csv.Error,UnicodeDecodeError):return [],[error(1,'file','Invalid UTF-8 CSV')]
        else:
            from openpyxl import load_workbook
            try:
                with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                    if sum(info.file_size for info in archive.infolist())>MAX_BYTES*10:raise ValueError('expanded workbook exceeds size limit')
                wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
                ws=wb['Examples'] if 'Examples' in wb.sheetnames else wb.active
                if ws.max_row>MAX_ROWS+1 or ws.max_column>1000:raise ValueError('workbook exceeds row or column limit')
                table=[]
                for n,line in enumerate(ws.iter_rows(),1):
                    table.append([cell.value for cell in line])
                    for cell in line:
                        if cell.data_type=='f':errors.append(error(n-1,cell.coordinate,'Spreadsheet formulas are not permitted'))
                wb.close()
            except ValueError:raise
            except Exception:return [],[error(1,'file','Invalid Excel workbook')]
        if not table:return [],[error(1,'file','File has no header')]
        header=table[0]
        if any(not isinstance(key,str) or not key for key in header) or len(set(header))!=len(header):
            return [],[error(1,'file','Headers must be distinct nonempty strings')]
        allowed=columns(schema,classifiers)
        for key in header:
            if key not in allowed:errors.append(error(1,key,'Unknown column'))
        rows=[]
        for n,line in enumerate(table[1:],1):
            if all(value in (None,'') for value in line):continue
            if len(line)!=len(header):errors.append(error(n,'file','Row length differs from header'))
            row={'values':{},'labels':{}}
            for key,value in zip(header,line):
                value='' if value is None else value
                if isinstance(value,(date,datetime)):value=value.isoformat()
                if key in ('id','occurred_at'):row[key]=str(value)
                elif key=='values':
                    try:row[key]=json.loads(value)
                    except (ValueError,TypeError):errors.append(error(n,key,'Expected JSON object'))
                elif key.startswith('values.'):
                    field=key[7:];field_schema=schema.get('properties',{}).get(field,{})
                    if value=='' and field not in schema.get('required',[]):continue
                    if field_schema.get('type')!='string' and isinstance(value,str):
                        try:value=json.loads(value)
                        except ValueError:errors.append(error(n,key,'Expected JSON value matching the schema'))
                    row['values'][field]=value
                elif '.' in key:
                    cid,field=key.rsplit('.',1)
                    if cid in {c['id'] for c in classifiers}:row['labels'].setdefault(cid,{})[field]=value
            rows.append(row)
    else:raise ValueError('supported formats: csv, json, jsonl, xlsx')
    if not 1<=len(rows)<=MAX_ROWS:raise ValueError('preview between one and 200 rows')
    return rows,errors


def validate_rows(rows,schema,classifiers,*,require_labels=True):
    validator=Draft202012Validator(checked_schema(schema));errors=[];valid=[];seen=set()
    for n,row in enumerate(rows,1):
        start=len(errors)
        if not isinstance(row,dict):errors.append(error(n,'row','Expected a structured item row'));continue
        identifier=row.get('id')
        if not isinstance(identifier,str) or not identifier.strip():errors.append(error(n,'id','Item identity required'))
        elif identifier in seen:errors.append(error(n,'id','Duplicate item identity'))
        else:seen.add(identifier)
        try:occurred=timestamp(row['occurred_at'])
        except (ValueError,KeyError,AttributeError,TypeError):occurred=None;errors.append(error(n,'occurred_at','Valid ISO 8601 item date required'))
        values=row.get('values')
        for failure in validator.iter_errors(values):errors.append(error(n,'values',failure.message))
        if isinstance(values,dict):
            for classifier in classifiers:
                config=classifier['config']
                try:DecisionTask(classifier['id'],tuple(c['label'] for c in config['classes']),config['question'],config.get('input_field','text')).validate_target(Item(identifier,values))
                except ValueError as exc:errors.append(error(n,'values',str(exc)))
        labels=row.get('labels',{})
        if not isinstance(labels,dict):labels={};errors.append(error(n,'labels','Labels must be an object'))
        labels=dict(labels)
        if set(labels)-{c['id'] for c in classifiers}:errors.append(error(n,'labels','Unknown classifier labels'))
        if require_labels:
            for classifier in classifiers:
                cid=classifier['id'];data=labels.get(cid,{})
                data=dict(data) if isinstance(data,dict) else {}
                if data.get('rationale') is None:data['rationale']=''
                labels[cid]=data
                if data.get('label') not in [c['label'] for c in classifier['config']['classes']]:errors.append(error(n,cid+'.label','Choose an explicit configured label'))
                if not isinstance(data.get('rationale'),str):errors.append(error(n,cid+'.rationale','Explanation must be text when supplied'))
        try:
            item=json.loads(encoded({k:v for k,v in row.items() if k not in ('labels','model_representation')}))
            item['occurred_at']=occurred
            if len(encoded(item).encode())>MAX_BYTES:raise ValueError('item exceeds size limit')
        except (ValueError,TypeError):errors.append(error(n,'values','Item is not bounded finite JSON'));item=None
        if len(errors)==start:
            valid.append({'row':n,'item':item,'labels':labels,'model_representation':{c['id']:values for c in classifiers}})
    return valid,errors
