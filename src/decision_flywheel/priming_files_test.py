"""Synthetic structured imports: errors never silently discard rows or labels."""
import base64
import csv
import io
import json
import pytest
from .priming_files import read_rows, template, validate_rows

SCHEMA={'type':'object','properties':{'title':{'type':'string','minLength':1},'abstract':{'type':'string','minLength':1},'count':{'type':'integer'}},'required':['title','abstract'],'additionalProperties':False}
CLASSIFIERS=[{'id':'editor','revision':2,'name':'Editorial','config':{'question':'Publish?', 'input_field':'abstract','classes':[{'label':'publish'},{'label':'not publish'}]}}, {'id':'scope','revision':1,'name':'Scope','config':{'question':'In scope?','input_field':'abstract','classes':[{'label':'yes'},{'label':'no'}]}}]
ROW={'id':'synthetic-1','occurred_at':'2026-10-14','values':{'title':'Synthetic paper','abstract':'Synthetic evidence','count':3},'labels':{'editor':{'label':'publish','rationale':'Synthetic novelty'},'scope':{'label':'yes','rationale':'Synthetic fit'}}}

@pytest.mark.parametrize('kind',['json','jsonl','csv','xlsx'])
def test_all_formats_preserve_values_and_explicit_multiple_classifier_labels(kind):
    if kind in ('csv','xlsx'):
        tab={'id':ROW['id'],'occurred_at':ROW['occurred_at'],**{'values.'+k:v for k,v in ROW['values'].items()},**{f'{cid}.{k}':v for cid,data in ROW['labels'].items() for k,v in data.items()}}
        if kind=='csv':
            f=io.StringIO();writer=csv.DictWriter(f,fieldnames=list(tab));writer.writeheader();writer.writerow(tab);content=f.getvalue().encode()
        else:
            from openpyxl import Workbook
            wb=Workbook();ws=wb.active;ws.append(list(tab));ws.append(list(tab.values()));f=io.BytesIO();wb.save(f);content=f.getvalue()
    else:content=json.dumps([ROW] if kind=='json' else ROW).encode()
    rows,errors=read_rows(kind,base64.b64encode(content).decode(),SCHEMA,CLASSIFIERS)
    valid,validation=validate_rows(rows,SCHEMA,CLASSIFIERS)
    assert errors==validation==[]
    assert valid[0]['item']['values']==ROW['values']
    assert valid[0]['labels']==ROW['labels']
    assert valid[0]['model_representation']=={'editor':ROW['values'],'scope':ROW['values']}


def test_templates_have_schema_fields_label_choices_and_no_invented_examples():
    for kind in ('csv','xlsx'):
        result=template(kind,SCHEMA,CLASSIFIERS)
        if kind=='csv':assert 'values.abstract,values.count,editor.label,editor.rationale,scope.label,scope.rationale' in result['content']
        else:
            from openpyxl import load_workbook
            wb=load_workbook(io.BytesIO(base64.b64decode(result['content'])))
            assert wb['Examples'].max_row==1
            assert len(wb['Examples'].data_validations.dataValidation)==2
        assert result['classifiers'][0]['labels']==['publish','not publish']


def test_row_errors_identify_missing_fields_and_unknown_labels():
    broken={**ROW,'values':{'title':'Paper'},'labels':{'editor':{'label':'invented','rationale':' '}}}
    valid,errors=validate_rows([ROW,broken],SCHEMA,CLASSIFIERS)
    assert len(valid)==1
    assert all(error['row']==2 for error in errors)
    assert {error['field'] for error in errors}>={'values','editor.label','scope.label'}


def test_malformed_jsonl_and_duplicate_ids_are_not_silently_accepted():
    rows,errors=read_rows('jsonl',base64.b64encode((json.dumps(ROW)+'\n{broken').encode()).decode(),SCHEMA,CLASSIFIERS)
    assert errors[0]['row']==2
    _,errors=validate_rows([ROW,ROW],SCHEMA,CLASSIFIERS)
    assert any(error['field']=='id' for error in errors)


def test_spreadsheet_formulas_and_unknown_columns_are_rejected():
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.append(['id','occurred_at','values.title','unknown']);ws.append(['a','2026-10-14','=WEBSERVICE("https://example.invalid")','x'])
    f=io.BytesIO();wb.save(f)
    _,errors=read_rows('xlsx',base64.b64encode(f.getvalue()).decode(),SCHEMA,CLASSIFIERS)
    assert any('formula' in error['message'].lower() for error in errors)
    assert any(error['field']=='unknown' for error in errors)


@pytest.mark.parametrize('vote',[{'label':'publish'},{'label':'publish','rationale':''},{'label':'publish','rationale':None}])
def test_explicit_label_accepts_optional_explanation_without_inventing_a_reason(vote):
    row={**ROW,'labels':{'editor':vote}}
    valid,errors=validate_rows([row],SCHEMA,CLASSIFIERS[:1])
    assert errors==[] and valid[0]['labels']['editor']=={'label':'publish','rationale':''}
